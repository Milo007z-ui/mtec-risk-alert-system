import os
"""ขั้นที่ 3 (ต่อ): เติม V จากค่าที่อธิบดี ทช. เปิดเผย ให้ 31 วงที่ข้อมูลเปิดของ ทช. ไม่มี + ชีตตรวจเทียบ กทม."""
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
P = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "outputs", "ขั้นที่3-AADT-ทางหลวงชนบท-3ปี.xlsx")
SRC = {
 "BH": (90303, "อธิบดีกรมทางหลวงชนบท (นายพิชิต หุ่นศิริ) 16 มี.ค. 2569: \"โดยมีปริมาณจราจรมากถึง 90,303 คันต่อวัน\" (สะพานภูมิพล 1 และ 2)",
        "https://mgronline.com/business/detail/9690000025473", "https://www.fm91bkk.com/newsarticle/68309"),
 "RP": (120000, "อธิบดีกรมทางหลวงชนบท (นายมนตรี เดชาสกุลสม) 20 มิ.ย. 2568: \"ถนนราชพฤกษ์ไม่สามารถรองรับการจราจรที่มีปริมาณหนาแน่นมากถึง 120,000 คันต่อวัน\" (โครงการสะพานข้ามคลองมหาสวัสดิ์ กม. 12+850–15+100 อยู่ในตอน กม. 11–28)",
        "https://mgronline.com/business/detail/9680000058163", "https://www.bangkokbiznews.com/business/economic/1185869"),
}
wb = load_workbook(P); ws = wb["38 คลัสเตอร์ + Accident Rate"]
h = {c.value: c.column for c in ws[1]}
for name in ("แหล่ง V (ค่าที่ ทช. เปิดเผย)", "ลิงก์อ้างอิง V ①", "ลิงก์อ้างอิง V ②"):
    if name not in h: ws.cell(1, ws.max_column + 1, name).font = Font(bold=True, color="FFFFFF"); ws.cell(1, ws.max_column, name).fill = PatternFill("solid", fgColor="1C5CAB")
h = {c.value: c.column for c in ws[1]}; L = lambda x: ws.cell(1, h[x]).column_letter
n = 0
for r in range(2, ws.max_row + 1):
    if ws.cell(r, h["ผล"]).value == "ครบ 3 ปี": continue
    code = ws.cell(r, h["รหัสสายทาง ทช."]).value
    v, txt, u1, u2 = SRC["BH" if code in ("ส.012", "ส.013") else "RP"]
    ws.cell(r, h["ผล"]).value = "ใช้ค่าที่อธิบดี ทช. เปิดเผย"
    ws.cell(r, h["ผล"]).fill = PatternFill("solid", fgColor="FFF2CC")
    ws[f"{L('V = AADT เฉลี่ย 3 ปี')}{r}"] = v
    ws[f"{L('ตัวตั้ง A×1,000,000')}{r}"] = f"={L('A')}{r}*1000000"
    ws[f"{L('ตัวหาร 365×T×V×L')}{r}"] = f"=365*{L('T (ปี)')}{r}*{L('V = AADT เฉลี่ย 3 ปี')}{r}*{L('L (กม.)')}{r}"
    ws[f"{L('R (ครั้ง/ล้านคัน-กม.)')}{r}"] = f"={L('ตัวตั้ง A×1,000,000')}{r}/{L('ตัวหาร 365×T×V×L')}{r}"
    for x in ("V = AADT เฉลี่ย 3 ปี", "ตัวตั้ง A×1,000,000", "ตัวหาร 365×T×V×L"): ws[f"{L(x)}{r}"].number_format = "#,##0"
    ws[f"{L('R (ครั้ง/ล้านคัน-กม.)')}{r}"].number_format = "0.0000"
    ws.cell(r, h["แหล่ง V (ค่าที่ ทช. เปิดเผย)"]).value = txt
    for col, u in (("ลิงก์อ้างอิง V ①", u1), ("ลิงก์อ้างอิง V ②", u2)):
        c = ws.cell(r, h[col]); c.value = u; c.hyperlink = u; c.font = Font(color="0563C1", underline="single")
    n += 1
ws.column_dimensions[L("แหล่ง V (ค่าที่ ทช. เปิดเผย)")].width = 70
for x in ("ลิงก์อ้างอิง V ①", "ลิงก์อ้างอิง V ②"): ws.column_dimensions[L(x)].width = 45
# ชีตตรวจเทียบ กทม.
if "ตรวจเทียบ กทม. Bridge67" in wb.sheetnames: del wb["ตรวจเทียบ กทม. Bridge67"]
w2 = wb.create_sheet("ตรวจเทียบ กทม. Bridge67")
for row in [("รายการ", "ค่า", "ที่มา"),
            ("สะพานภูมิพล 1 นับ 12 ชม. (7.00–19.00 น.) 16 ต.ค. 2567 · มุ่งหน้าพระราม 3", 45712, "https://traffic.bangkok.go.th/re_intersection/intersection/2567/Bridge67.pdf"),
            ("สะพานภูมิพล 1 นับ 12 ชม. · มุ่งหน้าสุขสวัสดิ์", 29249, "https://traffic.bangkok.go.th/re_intersection/intersection/2567/Bridge67.pdf"),
            ("สะพานภูมิพล 1 รวม 12 ชม. (สำนักการจราจรและขนส่ง กทม.)", 74961, "https://traffic.bangkok.go.th/re_intersection/intersection/2567/Bridge67.pdf"),
            ("ค่าทั้งวันที่อธิบดี ทช. เปิดเผย (16 มี.ค. 2569)", 90303, "https://mgronline.com/business/detail/9690000025473"),
            ("ส่วนต่าง = รถช่วงกลางคืน 19.00–7.00 น. (โดยประมาณ)", "=B5-B4", ""),
            ("สัดส่วนกลางคืนต่อทั้งวัน", "=B6/B5", "ประมาณ 17% สอดคล้องกับลักษณะถนนเมืองทั่วไป → ค่าของ ทช. สมเหตุสมผล")]:
    w2.append(list(row))
for c in w2[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1C5CAB")
for r in range(2, 7): w2.cell(r, 2).number_format = "#,##0"
w2.cell(7, 2).number_format = "0.0%"
for r in range(2, 6): c = w2.cell(r, 3); c.hyperlink = c.value; c.font = Font(color="0563C1", underline="single")
w2.column_dimensions["A"].width = 70; w2.column_dimensions["B"].width = 14; w2.column_dimensions["C"].width = 80
wb.save(P); print("updated", n, "rows")
