import os
"""ขั้นที่ 3: ทางหลวงชนบท (ทช.) 38 คลัสเตอร์ · AADT ปี 2566–2568 จาก dataportal.drr.go.th ชุด aadt1 (sum_AADT > 0 เท่านั้น)"""
import numpy as np, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
PROJ = os.environ.get("MTEC_OUT", os.path.join(HERE, "..", "outputs")) + "/"
DL = os.environ.get("MTEC_RAW", os.path.join(HERE, "..", "data-raw")) + "/"
DS = "https://dataportal.drr.go.th/dataset/3833590c-a0f4-4c99-baf1-8c540aadebb5"
RES = {2566: ("pcu-update-66.csv", DS + "/resource/3ae42fe1-d7e4-4784-a257-81475f5b8243", "ปริมาณจราจร 2566"),
       2567: ("pcu-update-67.csv", DS + "/resource/d0675c68-510b-45e1-b865-1ce261814948", "ปริมาณจราจร 2567"),
       2568: ("2568_traffic_aadt_drr.csv", DS + "/resource/0d612f16-e188-4604-b61c-c0da07854ef9", "2568_traffic_aadt_drr")}
Y = (2566, 2567, 2568)
D = {}
for y, (f, _, _) in RES.items():
    d = pd.read_csv(DL + f, encoding="utf-8-sig"); d["c"] = d.road_code.astype(str).str.replace(" ", ""); D[y] = d

k = pd.read_excel(PROJ + "AADT-ครบทุกคลัสเตอร์-630-3ปี.xlsx", sheet_name="คลัสเตอร์630")
t = k[k.หน่วยงาน == "ทช."].copy()
def seg(code, km_txt):          # นบ.3021 แบ่งตอนตาม กม. ในชื่อ road_name
    if code != "นบ.3021": return None
    a = float(str(km_txt).split("–")[0])
    return "ถนนราชพฤกษ์ (กม.ที่ 0+000 - 11+000)" if a < 11 else "ถนนราชพฤกษ์ (กม.ที่ 11+000 - 28+000)" if a < 28 else "ถนนราชพฤกษ์ (กม.ที่ 28+000 - 42+200)"
rows = []
for r in t.itertuples():
    code = r.route_code; s = seg(code, getattr(r, "_6"))
    o = {"id": r.id, "ถนน": r.ถนน, "จังหวัด": r.จังหวัด, "รหัสสายทาง ทช.": code, "กม. คลัสเตอร์": getattr(r, "_6"),
         "แถวที่ใช้ใน ทช. (road_name)": s or "", "A": r.A, "เสียชีวิต": r.เสียชีวิต, "บาดเจ็บ": r.บาดเจ็บ, "SI": r.SI, "ระดับ_SI": r.ระดับ_SI}
    for y in Y:
        d = D[y]; m = d[d.c == code]
        if s: m = m[m.road_name == s]
        v = m.sum_AADT.max() if len(m) else np.nan
        o[f"AADT {y}"] = v if v and v > 0 else np.nan
        o[f"สถานะ {y}"] = "มีค่า" if v and v > 0 else ("ทช. ให้ค่า 0 (ไม่ได้สำรวจ)" if len(m) else "ไม่มีแถวนี้ในไฟล์")
        o[f"ลิงก์ {y}"] = RES[y][1]
        o[f"วิธีค้น {y}"] = f"ค้น road_code = {code}" + (f" และ road_name = {s}" if s else "") + " → คอลัมน์ sum_AADT"
    n = sum(pd.notna(o[f"AADT {y}"]) for y in Y)
    o["ผล"] = "ครบ 3 ปี" if n == 3 else ("ไม่มีข้อมูลจาก ทช. ทั้ง 3 ปี" if n == 0 else f"มี {n} ปี")
    rows.append(o)
C = pd.DataFrame(rows).sort_values(["ผล", "รหัสสายทาง ทช.", "id"]).reset_index(drop=True)
C["V = AADT เฉลี่ย 3 ปี"] = None; C["T (ปี)"] = 3; C["L (กม.)"] = 1
print(C.groupby(["ผล", "รหัสสายทาง ทช."]).agg(n=("id", "size"), A=("A", "sum")).to_string())
print(C.groupby("ผล").agg(n=("id", "size"), A=("A", "sum")).to_string())

codes = sorted(set(C["รหัสสายทาง ทช."]))
raw = pd.concat([D[y][D[y].c.isin(codes)].assign(ไฟล์=RES[y][0]) for y in Y])[["ไฟล์", "budget_year", "road_code", "road_name", "sum_AADT"]]
summary = (C.groupby(["รหัสสายทาง ทช.", "แถวที่ใช้ใน ทช. (road_name)"]).agg(ชื่อถนน=("ถนน", lambda s: s.value_counts().index[0]), คลัสเตอร์=("id", "size"),
           อุบัติเหตุ=("A", "sum"), **{f"AADT {y}": (f"AADT {y}", "first") for y in Y}, ผล=("ผล", "first")).reset_index().sort_values(["ผล", "อุบัติเหตุ"], ascending=[True, False]))

out = PROJ + "ขั้นที่3-AADT-ทางหลวงชนบท-3ปี.xlsx"
with pd.ExcelWriter(out, engine="openpyxl") as w:
    summary.to_excel(w, sheet_name="สรุปรายสายทาง", index=False)
    C.to_excel(w, sheet_name="38 คลัสเตอร์ + Accident Rate", index=False)
    raw.to_excel(w, sheet_name="แถวดิบจากไฟล์ ทช.", index=False)
    pd.DataFrame({"รายการ": ["หน้าชุดข้อมูล", "ปี 2566", "ปี 2567", "ปี 2568", "ปี 2569 (ไม่ใช้ในสูตร)", "หลักการ", "ตรวจแล้ว", "สูตร"],
                  "รายละเอียด": [
        f"กรมทางหลวงชนบท · ชุดข้อมูล aadt1 'สถิติปริมาณจราจร': {DS}",
        f"{RES[2566][2]} ({RES[2566][0]}): {RES[2566][1]}",
        f"{RES[2567][2]} ({RES[2567][0]}): {RES[2567][1]}",
        f"{RES[2568][2]} ({RES[2568][0]}): {RES[2568][1]}",
        "2569_traffic_aadt_drr.csv: ไม่มีแถว ส.012 / ส.013 และราชพฤกษ์มีแค่ตอน กม. 28–42.2 เหมือนปี 2568",
        "ใช้ค่า sum_AADT ของแถวที่ road_code ตรงกับรหัสสายทางในข้อมูลอุบัติเหตุ · ค่า 0 ถือว่าไม่ได้สำรวจ (ไม่ใช่ไม่มีรถ) จึงไม่นำมาใช้ · นบ.3021 แบ่ง 3 ตอนตาม กม. ใช้ตอนที่ กม. คลัสเตอร์อยู่",
        "ค่าในไฟล์ที่ดาวน์โหลด = ไฟล์บนหน้าเว็บ (รหัส resource ตรงกับชื่อไฟล์ใน API ของ dataportal.drr.go.th)",
        "R = A × 1,000,000 / (365 × T × V × L) · V = AADT เฉลี่ย 2566–2568 · T = 3 ปี · L = 1 กม."]}).to_excel(w, sheet_name="แหล่งข้อมูลและวิธีตรวจ", index=False)

wb = load_workbook(out); blue = PatternFill("solid", fgColor="1C5CAB"); orange = PatternFill("solid", fgColor="EB6834"); red = PatternFill("solid", fgColor="F8CBAD")
for ws in wb:
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = blue; c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 42; ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
    for j in range(1, ws.max_column + 1): ws.column_dimensions[get_column_letter(j)].width = 14
ws = wb["38 คลัสเตอร์ + Accident Rate"]; h = {c.value: c.column for c in ws[1]}; n = ws.max_row
extra = ["ตัวตั้ง A×1,000,000", "ตัวหาร 365×T×V×L", "R (ครั้ง/ล้านคัน-กม.)"]
for x in extra: ws.cell(1, ws.max_column + 1, x)
h = {c.value: c.column for c in ws[1]}; L = lambda x: get_column_letter(h[x])
for x in extra + ["V = AADT เฉลี่ย 3 ปี"]: ws.cell(1, h[x]).fill = orange; ws.cell(1, h[x]).font = Font(bold=True, color="FFFFFF")
for r in range(2, n + 1):
    for y in Y:
        c = ws.cell(r, h[f"ลิงก์ {y}"]); c.hyperlink = c.value; c.value = "เปิดลิงก์"; c.font = Font(color="0563C1", underline="single")
        ws.cell(r, h[f"AADT {y}"]).number_format = "#,##0"
    if ws.cell(r, h["ผล"]).value != "ครบ 3 ปี":
        for x in ("ผล",) + tuple(f"สถานะ {y}" for y in Y): ws.cell(r, h[x]).fill = red
        continue
    ws[f"{L('V = AADT เฉลี่ย 3 ปี')}{r}"] = f"=AVERAGE({L('AADT 2566')}{r},{L('AADT 2567')}{r},{L('AADT 2568')}{r})"
    ws[f"{L(extra[0])}{r}"] = f"={L('A')}{r}*1000000"
    ws[f"{L(extra[1])}{r}"] = f"=365*{L('T (ปี)')}{r}*{L('V = AADT เฉลี่ย 3 ปี')}{r}*{L('L (กม.)')}{r}"
    ws[f"{L(extra[2])}{r}"] = f"={L(extra[0])}{r}/{L(extra[1])}{r}"
    for x in ("V = AADT เฉลี่ย 3 ปี", extra[0], extra[1]): ws[f"{L(x)}{r}"].number_format = "#,##0"
    ws[f"{L(extra[2])}{r}"].number_format = "0.0000"
for x, wd in (("ถนน", 34), ("แถวที่ใช้ใน ทช. (road_name)", 30), ("ผล", 26)): ws.column_dimensions[L(x)].width = wd
for y in Y: ws.column_dimensions[L(f"สถานะ {y}")].width = 22; ws.column_dimensions[L(f"วิธีค้น {y}")].width = 40
ws = wb["สรุปรายสายทาง"]; ws.column_dimensions["B"].width = 32; ws.column_dimensions["C"].width = 40; ws.column_dimensions["J"].width = 26
for r in range(2, ws.max_row + 1):
    for j in (6, 7, 8): ws.cell(r, j).number_format = "#,##0"
    if ws.cell(r, 9).value != "ครบ 3 ปี": ws.cell(r, 9).fill = red
ws = wb["แถวดิบจากไฟล์ ทช."]; ws.column_dimensions["A"].width = 26; ws.column_dimensions["D"].width = 52
ws = wb["แหล่งข้อมูลและวิธีตรวจ"]; ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 150
wb.save(out); print("saved", out)
print(summary.to_string())
