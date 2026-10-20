"""แบ่งระดับจุดเสี่ยง 598 คลัสเตอร์ (ข้อมูลเปิดของหน่วยงาน ครบ 3 ปี) : Accident Rate + RQC ร่วมกับ Accident Severity (SI)
ทล.  : สถานีตัวแทนที่ ทล. กำหนดในชั้น hgds_aadt_line_2567 → ค่าสถานีเดียวกัน 2566 (opendata aadt-66) / 2567 / 2568 (GeoServer = opendata)
กทพ. : ยอดผ่านทุกด่านในสาย ÷ 365 (ByPlaza) ปีงบ 2566–2568
ทช.  : sum_AADT รายสายทาง (dataportal.drr.go.th aadt1) 2566–2568 เฉพาะค่า > 0"""
import os
import json, re
import numpy as np, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
PROJ = os.environ.get("MTEC_OUT", os.path.join(HERE, "..", "outputs")) + "/"
DL = os.environ.get("MTEC_RAW", os.path.join(HERE, "..", "data-raw")) + "/"
SP = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data") + "/"
Y = (2566, 2567, 2568); T, L, K = 3, 1, 1.645

k = pd.read_excel(PROJ + "AADT-ครบทุกคลัสเตอร์-630-3ปี.xlsx", sheet_name="คลัสเตอร์630")
base = k[["id", "ถนน", "จังหวัด", "route_code", "หน่วยงาน", "กม. คลัสเตอร์", "ตอนที่จับคู่", "A", "เสียชีวิต", "บาดเจ็บ", "SI", "ระดับ_SI"]].copy()

# ---------- ทล.
J = json.load(open(SP + "doh70.json", encoding="utf-8"))
line = {(int(x["highway_code"]), int(x["road_code"])): x for x in J["line"]}
def uniq(rows):
    d = {}
    for x in rows:
        d.setdefault(x["station_id"], set()).add((x["highway_code"], x["km"], x["vehcat_aadt_all"]))
    out = {}
    for s, v in d.items():
        assert len({a for _, _, a in v}) == 1, (s, v); out[s] = sorted(v)[0]
    return out
p67, p68 = uniq(J["p67"]), uniq(J["p68"])
t66 = pd.read_excel(DL + "aadt-66.xlsx")
def v66(h, km):
    m = t66[(pd.to_numeric(t66["ทางหลวงสาย"], errors="coerce") == h) & (t66["จุดสำรวจ"].astype(str).str.strip() == km)]["รวม"]
    v = set(m); assert len(v) <= 1, (h, km, v); return int(m.iloc[0]) if len(m) else np.nan
doh = {}
for key, x in line.items():
    s = x["station_id"]; h, km, a67 = p67[s]
    assert a67 == x["vehcat_aadt_all"]
    a68 = p68.get(s, (None, None, np.nan))[2]
    doh[key] = {"station_id": s, "สถานี": f"ทล.{int(h)} กม.{km}", "AADT 2566": v66(int(h), km), "AADT 2567": a67, "AADT 2568": a68}
# ---------- กทพ.
ex = pd.concat([pd.DataFrame(json.load(open(SP + f"exat_{y}.json", encoding="utf-8"))["result"]) for y in Y])
ex["year"] = ex.year.astype(int); ex["traffic"] = ex.traffic.astype(int)
exat = (ex.groupby(["highway_name", "year"]).traffic.sum() / 365).unstack()
# ---------- ทช.
DRR = {2566: "pcu-update-66.csv", 2567: "pcu-update-67.csv", 2568: "2568_traffic_aadt_drr.csv"}
drr = {}
for y, f in DRR.items():
    d = pd.read_csv(DL + f, encoding="utf-8-sig"); d["c"] = d.road_code.astype(str).str.replace(" ", "")
    for c, g in d[d.sum_AADT > 0].groupby("c"):
        if c != "นบ.3021": drr.setdefault(c, {})[y] = g.sum_AADT.max()

rows = []
for r in base.itertuples(index=False):
    o = {"แหล่ง V": "", "สถานีตัวแทน/สายทาง": ""}
    if r.หน่วยงาน == "ทล.":
        m = re.search(r"ทางหลวง (\d+) ตอนควบคุม (\d+)", r.ตอนที่จับคู่); x = doh[(int(m.group(1)), int(m.group(2)))]
        o.update({f"AADT {y}": x[f"AADT {y}"] for y in Y}); o["แหล่ง V"] = "ทล. สถานีตัวแทนตอนควบคุม"
        o["สถานีตัวแทน/สายทาง"] = f"station_id {x['station_id']} ({x['สถานี']})"
    elif r.หน่วยงาน == "กทพ.":
        name = re.sub(r"^ทางพิเศษ(?=[^ก])", "", r.ตอนที่จับคู่)
        if name in exat.index:
            o.update({f"AADT {y}": exat.loc[name, y] for y in Y}); o["แหล่ง V"] = "กทพ. ยอดผ่านด่านทั้งสาย ÷ 365"; o["สถานีตัวแทน/สายทาง"] = name
    else:
        c = str(r.route_code).replace(" ", "")
        if c in drr and len(drr[c]) == 3:
            o.update({f"AADT {y}": drr[c][y] for y in Y}); o["แหล่ง V"] = "ทช. AADT รายสายทาง"; o["สถานีตัวแทน/สายทาง"] = c
    rows.append(o)
V = pd.DataFrame(rows)
for y in Y:
    if f"AADT {y}" not in V: V[f"AADT {y}"] = np.nan
D = pd.concat([base.reset_index(drop=True), V], axis=1)
D["ครบ3ปี"] = D[[f"AADT {y}" for y in Y]].notna().all(axis=1)
print("ครบ 3 ปี", D.ครบ3ปี.sum(), "| แยก:", D[D.ครบ3ปี].หน่วยงาน.value_counts().to_dict())
print("ทล. ไม่ครบ:", D[(D.หน่วยงาน == "ทล.") & ~D.ครบ3ปี][["id", "ตอนที่จับคู่", "สถานีตัวแทน/สายทาง", "AADT 2566", "AADT 2567", "AADT 2568"]].to_string())

M = D[D.ครบ3ปี].copy()
M["V"] = M[[f"AADT {y}" for y in Y]].mean(axis=1)
M["m"] = 365 * T * M.V * L / 1e6
M["R"] = M.A / M.m
for ag, g in M.groupby("หน่วยงาน"): M.loc[g.index, "Ra"] = g.A.sum() / g.m.sum()
M["Rc"] = M.Ra + K * np.sqrt(M.Ra / M.m) + 1 / (2 * M.m)
M["ระดับ_Rate"] = np.where(M.R > M.Rc, "high", np.where(M.R > M.Ra, "medium", "low"))
rk = {"low": 0, "medium": 1, "high": 2}; inv = {v: s for s, v in rk.items()}
M["ระดับรวม (SI+Rate)"] = [inv[max(rk[a], rk[b])] for a, b in zip(M.ระดับ_SI, M.ระดับ_Rate)]
M["เหตุผลที่ได้ระดับนี้"] = [("ทั้ง SI และ Rate" if a == b else ("SI" if rk[a] > rk[b] else "Rate")) for a, b in zip(M.ระดับ_SI, M.ระดับ_Rate)]
ct = pd.crosstab(M.ระดับ_SI, M.ระดับ_Rate).reindex(index=["high", "medium", "low"], columns=["high", "medium", "low"], fill_value=0)
print(ct); print("SI:", M.ระดับ_SI.value_counts().to_dict(), "| Rate:", M.ระดับ_Rate.value_counts().to_dict(), "| รวม:", M["ระดับรวม (SI+Rate)"].value_counts().to_dict())
print("Ra:", M.groupby("หน่วยงาน").Ra.first().round(4).to_dict(), "n:", M.หน่วยงาน.value_counts().to_dict())
X = D[~D.ครบ3ปี][["id", "ถนน", "จังหวัด", "route_code", "หน่วยงาน", "A", "SI", "ระดับ_SI"]].copy()
X["เหตุผลที่ไม่อยู่ใน 598"] = np.where(X.route_code.astype(str).str.startswith("ส.0"), "สะพานภูมิพล: ข้อมูลเปิด ทช. = 0 / ไม่มีแถว",
                                np.where(X.route_code.astype(str).str.startswith("นบ.3021"), "ราชพฤกษ์ กม.11–28: ข้อมูลเปิด ทช. = 0 / ไม่มีแถว",
                                         "E006: ไม่มีสายนี้ใน API กทพ."))
print("ไม่อยู่ใน 598:", len(X), X["เหตุผลที่ไม่อยู่ใน 598"].value_counts().to_dict())
M.to_pickle(SP + "M598.pkl"); X.to_pickle(SP + "X32.pkl"); ct.to_pickle(SP + "ct.pkl")
