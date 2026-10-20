import os
import numpy as np, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule

HERE = os.path.dirname(os.path.abspath(__file__))
PROJ = os.environ.get("MTEC_OUT", os.path.join(HERE, "..", "outputs")) + "/"
DL = os.environ.get("MTEC_RAW", os.path.join(HERE, "..", "data-raw")) + "/"
SP = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data") + "/"
M = pd.read_pickle(SP + "M598.pkl"); X = pd.read_pickle(SP + "X32.pkl"); ct = pd.read_pickle(SP + "ct.pkl")
order = {"high": 0, "medium": 1, "low": 2}
M = M.assign(_o=M["ระดับรวม (SI+Rate)"].map(order), _df=M.R / M.Rc).sort_values(["_o", "_df"], ascending=[True, False])
cols = ["id", "ถนน", "จังหวัด", "route_code", "หน่วยงาน", "กม. คลัสเตอร์", "สถานีตัวแทน/สายทาง", "แหล่ง V",
        "A", "เสียชีวิต", "บาดเจ็บ", "SI", "ระดับ_SI", "AADT 2566", "AADT 2567", "AADT 2568"]
out = PROJ + "ระดับจุดเสี่ยง-598คลัสเตอร์-SI+AccidentRate.xlsx"
with pd.ExcelWriter(out, engine="openpyxl") as w:
    M[cols].to_excel(w, sheet_name="ระดับจุดเสี่ยง 598", index=False)
    X.to_excel(w, sheet_name="32 วงที่ไม่มี V (ใช้ SI)", index=False)
    pd.DataFrame().to_excel(w, sheet_name="สรุปผล", index=False)
    pd.DataFrame({"ขั้น": [], "รายละเอียด": []}).to_excel(w, sheet_name="วิธีคำนวณและแหล่งข้อมูล", index=False)

wb = load_workbook(out)
blue = PatternFill("solid", fgColor="1C5CAB"); orange = PatternFill("solid", fgColor="EB6834")
LV = {"high": "F8CBAD", "medium": "FFE699", "low": "C6EFCE"}
ws = wb["ระดับจุดเสี่ยง 598"]; n = ws.max_row
calc = ["V = AADT เฉลี่ย 3 ปี", "T (ปี)", "L (กม.)", "MEV (ล้านคัน)", "R (ครั้ง/ล้านคัน)", "Ra ของหน่วยงาน", "Rc (K=1.645)", "R/Rc", "ระดับ_Rate", "ระดับรวม (SI+Rate)", "ได้ระดับนี้จาก"]
for c in calc: ws.cell(1, ws.max_column + 1, c)
h = {c.value: c.column for c in ws[1]}; Lc = lambda x: get_column_letter(h[x])
A, AG, SIL = Lc("A"), Lc("หน่วยงาน"), Lc("ระดับ_SI")
for r in range(2, n + 1):
    ws[f"{Lc(calc[0])}{r}"] = f"=AVERAGE({Lc('AADT 2566')}{r}:{Lc('AADT 2568')}{r})"
    ws[f"{Lc('T (ปี)')}{r}"] = 3; ws[f"{Lc('L (กม.)')}{r}"] = 1
    ws[f"{Lc(calc[3])}{r}"] = f"=365*{Lc('T (ปี)')}{r}*{Lc(calc[0])}{r}*{Lc('L (กม.)')}{r}/1000000"
    ws[f"{Lc(calc[4])}{r}"] = f"={A}{r}/{Lc(calc[3])}{r}"
    ws[f"{Lc(calc[5])}{r}"] = f"=SUMIF(${AG}$2:${AG}${n},{AG}{r},${A}$2:${A}${n})/SUMIF(${AG}$2:${AG}${n},{AG}{r},${Lc(calc[3])}$2:${Lc(calc[3])}${n})"
    ws[f"{Lc(calc[6])}{r}"] = f"={Lc(calc[5])}{r}+1.645*SQRT({Lc(calc[5])}{r}/{Lc(calc[3])}{r})+1/(2*{Lc(calc[3])}{r})"
    ws[f"{Lc(calc[7])}{r}"] = f"={Lc(calc[4])}{r}/{Lc(calc[6])}{r}"
    R_, Ra_, Rc_ = (f"{Lc(calc[i])}{r}" for i in (4, 5, 6))
    ws[f"{Lc(calc[8])}{r}"] = f'=IF({R_}>{Rc_},"high",IF({R_}>{Ra_},"medium","low"))'
    rs = f'IF({SIL}{r}="high",2,IF({SIL}{r}="medium",1,0))'; rr = f'IF({Lc(calc[8])}{r}="high",2,IF({Lc(calc[8])}{r}="medium",1,0))'
    ws[f"{Lc(calc[9])}{r}"] = f'=CHOOSE(MAX({rs},{rr})+1,"low","medium","high")'
    ws[f"{Lc(calc[10])}{r}"] = f'=IF({rs}={rr},"SI และ Rate เท่ากัน",IF({rs}>{rr},"SI","Rate"))'
    for x in ("AADT 2566", "AADT 2567", "AADT 2568", calc[0]): ws[f"{Lc(x)}{r}"].number_format = "#,##0"
    for x, f in ((calc[3], "#,##0.00"), (calc[4], "0.0000"), (calc[5], "0.0000"), (calc[6], "0.0000"), (calc[7], "0.00"), ("SI", "0.00")): ws[f"{Lc(x)}{r}"].number_format = f
for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = blue; c.alignment = Alignment(wrap_text=True, vertical="center")
for c in calc: ws.cell(1, h[c]).fill = orange
ws.row_dimensions[1].height = 45; ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
for j in range(1, ws.max_column + 1): ws.column_dimensions[get_column_letter(j)].width = 12
for x, wd in (("ถนน", 32), ("สถานีตัวแทน/สายทาง", 30), ("แหล่ง V", 26), ("ได้ระดับนี้จาก", 18)): ws.column_dimensions[Lc(x)].width = wd
for col in (SIL, Lc("ระดับ_Rate"), Lc("ระดับรวม (SI+Rate)")):
    for lv, color in LV.items():
        ws.conditional_formatting.add(f"{col}2:{col}{n}", FormulaRule(formula=[f'{col}2="{lv}"'], fill=PatternFill("solid", fgColor=color)))

# สรุปผล (ค่าคงที่จาก Python ตรวจซ้ำได้กับสูตรในชีตหลัก)
s = wb["สรุปผล"]; bold = Font(bold=True)
s.append(["① ตารางไขว้ ระดับ SI × ระดับ Accident Rate (RQC) — จำนวนคลัสเตอร์"]); s["A1"].font = Font(bold=True, size=12)
s.append(["SI \\ Rate", "high", "medium", "low", "รวม"])
for lv in ("high", "medium", "low"): s.append([lv, *[int(ct.loc[lv, c]) for c in ("high", "medium", "low")], int(ct.loc[lv].sum())])
s.append(["รวม", *[int(ct[c].sum()) for c in ("high", "medium", "low")], int(ct.values.sum())])
s.append([]); s.append(["② จำนวนคลัสเตอร์ตามวิธี"]); s.cell(s.max_row, 1).font = Font(bold=True, size=12)
s.append(["ระดับ", "SI อย่างเดียว", "Rate (RQC) อย่างเดียว", "SI + Rate (ใช้ระดับที่สูงกว่า)"])
for lv in ("high", "medium", "low"):
    s.append([lv, int((M.ระดับ_SI == lv).sum()), int((M.ระดับ_Rate == lv).sum()), int((M["ระดับรวม (SI+Rate)"] == lv).sum())])
s.append([]); s.append(["③ แยกตามหน่วยงาน"]); s.cell(s.max_row, 1).font = Font(bold=True, size=12)
s.append(["หน่วยงาน", "คลัสเตอร์", "อุบัติเหตุ A", "Ra (ครั้ง/ล้านคัน)", "SI+Rate: high", "medium", "low"])
for ag in ("ทล.", "กทพ.", "ทช."):
    g = M[M.หน่วยงาน == ag]
    s.append([ag, len(g), int(g.A.sum()), round(float(g.Ra.iloc[0]), 4), *[int((g["ระดับรวม (SI+Rate)"] == lv).sum()) for lv in ("high", "medium", "low")]])
s.append([]); s.append(["④ 15 อันดับ R/Rc สูงสุด (เกินค่าวิกฤติมากที่สุด)"]); s.cell(s.max_row, 1).font = Font(bold=True, size=12)
s.append(["id", "ถนน", "หน่วยงาน", "A", "V (คัน/วัน)", "R", "Rc", "R/Rc", "SI", "ระดับ_SI", "ระดับรวม"])
for r in M.assign(df=M.R / M.Rc).nlargest(15, "df").itertuples():
    s.append([r.id, r.ถนน, r.หน่วยงาน, int(r.A), round(r.V), round(r.R, 4), round(r.Rc, 4), round(r.df, 2), round(r.SI, 2), r.ระดับ_SI, M.loc[r.Index, "ระดับรวม (SI+Rate)"]])
for row in s.iter_rows():
    for c in row:
        if c.value in ("SI \\ Rate", "ระดับ", "หน่วยงาน", "id"):
            for cc in row: cc.font = Font(bold=True, color="FFFFFF"); cc.fill = blue
s.column_dimensions["A"].width = 14; s.column_dimensions["B"].width = 34
for j in "CDEFGHIJK": s.column_dimensions[j].width = 16

# 32 วง
x = wb["32 วงที่ไม่มี V (ใช้ SI)"]
for c in x[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = blue
x.column_dimensions["B"].width = 40; x.column_dimensions["I"].width = 48

# วิธีคำนวณ
v = wb["วิธีคำนวณและแหล่งข้อมูล"]; v.delete_rows(1, v.max_row)
for row in [
    ("ขั้น", "รายละเอียด"),
    ("1. ปริมาณจราจร V", "V = ค่าเฉลี่ยของ AADT ปี 2566, 2567, 2568 (คัน/วัน) — ใช้เฉพาะคลัสเตอร์ที่มีข้อมูลเปิดของหน่วยงานครบ 3 ปี = 598 คลัสเตอร์"),
    ("   ทล. (466)", "สถานีตัวแทน (station_id) ที่กรมทางหลวงกำหนดให้ตอนควบคุมในชั้นข้อมูล hgds_aadt_line_2567 → ค่าของสถานีเดียวกัน: 2566 จาก opendata aadt-66, 2567 และ 2568 จาก GeoServer (= opendata aadt-67/68 = AADT_Pnt_2568) · https://onemap-layers.doh.go.th/geoserver/wfs · https://opendata.doh.go.th/dataset/ed101df4-f0a1-4d7f-b9d9-76859b2ca73e"),
    ("   กทพ. (125)", "ยอดรถผ่านทุกด่านในสายทาง ÷ 365 (สถิติปริมาณจราจรจำแนกตามด่าน) ปีงบ 2566–2568 · https://exat-man.web.app/api/EXAT_TrafficStatByPlaza/{ปีงบ} · https://catalog.exat.co.th/dataset/traffic-stat"),
    ("   ทช. (7)", "sum_AADT รายสายทาง · https://dataportal.drr.go.th/dataset/3833590c-a0f4-4c99-baf1-8c540aadebb5"),
    ("2. Accident Rate (สมการ 2.2)", "R = A × 1,000,000 / (365 × T × V × L) = A / MEV  · A = จำนวนอุบัติเหตุ 3 ปี · T = 3 ปี · L = 1 กม. (คงที่) · MEV = 365×T×V×L/10⁶ (ล้านคันที่ผ่าน, HSM สมการ 4-8)"),
    ("3. Rate Quality Control", "Ra = ΣA / ΣMEV ของหน่วยงานเดียวกัน = ค่าเฉลี่ยถ่วงน้ำหนักด้วยปริมาณจราจร (AASHTO 2010 HSM สมการ 4-10 หน้า 4-43) · Rc = Ra + K√(Ra/MEV) + 1/(2·MEV) (ณัฐพงศ์ 2558 สมการ 3.4 หน้า 35 · HSM สมการ 4-11 หน้า 4-44) · DF = R/Rc (สมการ 3.3) · K = 1.645 ความเชื่อมั่น 95% (ณัฐพงศ์ หน้า 35 · HSM Exhibit 4-46 หน้า 4-41)"),
    ("4. ระดับตาม Rate", "ตามขั้นตอนของ MnDOT (2010) สไลด์ 3-18: เทียบกับค่าเฉลี่ยก่อน แล้วทดสอบค่าวิกฤติเฉพาะจุดที่สูงกว่าค่าเฉลี่ย → low: R ≤ Ra (ไม่เกินค่าเฉลี่ย) · medium: Ra < R ≤ Rc (สูงกว่าค่าเฉลี่ยแต่ไม่มีนัยสำคัญทางสถิติ — ชื่อระดับเป็นการตีความของผู้วิจัย) · high: R > Rc หรือ DF > 1 (ณัฐพงศ์ 2558 รูป 4.4 หน้า 54 · HSM หน้า 4-12, 4-41)"),
    ("5. ระดับตาม SI", "SI = (F + PI) / Total Accident (ณัฐพงศ์ 2558 สมการ 2.6 หน้า 7) · low: SI < 1 · medium: 1 ≤ SI < 2 · high: SI ≥ 2 · จุดตัด 1 และ 2 กำหนดโดยผู้วิจัย สอดคล้องกับช่วงที่ราบของการกระจาย SI (P76–P83 มีค่า SI = 1.000)"),
    ("6. รวม SI + Rate", "ระดับรวม = ระดับที่สูงกว่าระหว่าง ระดับ_SI กับ ระดับ_Rate → จุดที่เกินเกณฑ์ด้านอัตราหรือด้านความรุนแรงอย่างใดอย่างหนึ่งถือว่าเสี่ยง (WisDOT Figure 3 หน้า 9: Crash Rate > UCL OR KAB Crash Rate > KAB UCL · MnDOT สไลด์ 3-18: 'Crash rate may be low, but crash severity may be high' · ณัฐพงศ์ หน้า 55 เสนอให้ใช้ Accident Severity ร่วมกับวิธีอื่น)"),
    ("ทำไมแยก Ra ตามหน่วยงาน", "ปริมาณจราจรของแต่ละหน่วยงานวัดคนละแบบ (ทางพิเศษใช้ยอดผ่านด่านทั้งสาย ซึ่งสูงกว่า AADT รายจุดมาก) ถ้าใช้ Ra เดียวกันทั้งชุด ทางพิเศษจะถูกจัดเป็นความเสี่ยงต่ำเกือบทั้งหมด"),
    ("ข้อจำกัด", "ทช. มีเพียง 7 คลัสเตอร์ ค่า Ra ของกลุ่มจึงไม่เสถียร · 32 คลัสเตอร์ที่ไม่มีข้อมูลเปิด (สะพานภูมิพล 29, ราชพฤกษ์ กม.11–28 อีก 2, E006 อีก 1) ใช้ระดับจาก SI อย่างเดียว (ชีต 32 วง)"),
]: v.append(list(row))
for c in v[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = blue
v.column_dimensions["A"].width = 28; v.column_dimensions["B"].width = 160
for row in v.iter_rows(min_row=2):
    row[1].alignment = Alignment(wrap_text=True, vertical="top")

# อ้างอิงรายเกณฑ์
ref = wb.create_sheet("อ้างอิงเกณฑ์")
ref.append(["เกณฑ์ / สูตร", "แหล่งอ้างอิง", "ตำแหน่งในเอกสาร", "ข้อความสำคัญ", "ลิงก์"])
NATT = "ณัฐพงศ์ ซื่อสัตย์ (2558) วิทยานิพนธ์ มทส. การจัดทำแผนที่ความเสี่ยงของทางหลวงในประเทศไทย"
HSM = "https://ftp.granit.unh.edu/submissions/NHDOTOUT/craig/HSM/Chapter%204_NetworkScreening.pdf"
MN = "https://www.dot.state.mn.us/context-sensitive-solutions/pdf/2010/3-risk.pdf"
WI = "https://wisconsindot.gov/dtsdManuals/traffic-ops/manuals-and-standards/safety/crashrates/rates20.pdf"
for r in [
    ("Accident Rate R", NATT, "สมการ 2.2 หน้า 5 · สมการ 3.2 หน้า 35", "R = (A × 1,000,000) / (365 × T × V × L)", ""),
    ("Ra (สูตรคำนวณ)", "AASHTO (2010) Highway Safety Manual Ch.4 (ฉบับร่าง 2009)", "สมการ 4-10 หน้า 4-43 (PDF หน้า 49)", "weighted average crash rate for each population weighted by traffic volume", HSM),
    ("Ra (นิยาม)", NATT, "หน้า 7 และหน้า 35", "อัตราการเกิดอุบัติเหตุเฉลี่ยของทุก ๆ ส่วน บนถนน", ""),
    ("Rc", NATT, "สมการ 2.5 หน้า 6 · สมการ 3.4 หน้า 35", "Rc = Ra + K[Ra / (365TVL/10⁶)]^0.5 + 1 / (2 × 365TVL/10⁶)", ""),
    ("Rc (สูตรเดียวกันใน HSM)", "AASHTO (2010) HSM Ch.4 (ฉบับร่าง 2009)", "สมการ 4-11 หน้า 4-44 (PDF หน้า 50) · STEP 5 หน้าเดียวกัน", "Rc,i = Ra + P × [Ra / MEVi]^0.5 + 1 / (2 × MEVi) · Any intersection with an observed crash rate greater than the corresponding critical crash rate is flagged for further review", HSM),
    ("K = 1.645", "AASHTO (2010) HSM Ch.4 · " + NATT, "Exhibit 4-46 หน้า 4-41 · หน้า 35", "95 Percent → 1.645", HSM),
    ("DF = R/Rc และ DF > 1 = จุดเสี่ยง", NATT, "สมการ 3.3 หน้า 35 · รูป 4.4 หน้า 54", "(ค่า DANGEROUS FACTOR >1)", ""),
    ("R > Rc = flagged", "AASHTO (2010) HSM Ch.4", "หน้า 4-12 และ 4-41", "Sites that exceed their respective critical rate are flagged for further review", HSM),
    ("ลำดับขั้น: เทียบค่าเฉลี่ย แล้วทดสอบค่าวิกฤติ (ที่มาของ 3 ระดับ Rate)", "Minnesota DOT (2010) Risk Management and Safety, Session 3", "สไลด์ 3-18 (ตัวอย่างสไลด์ 3-21)", "Critical Crash Rate – Performed if the crash rate is ABOVE average. Tests the crash rate to see if there is statistical significance", MN),
    ("แยกกลุ่มอ้างอิงตามประเภทถนน (Ra แยกหน่วยงาน)", "AASHTO (2010) HSM Ch.4", "หน้า 4-43 STEP 3", "Divide the network into reference populations based on operational or geometric differences", HSM),
    ("Accident Severity SI", NATT, "สมการ 2.6 หน้า 7", "SI = (F + PI) / Total Accidents", ""),
    ("รวม Rate กับ Severity (อย่างใดอย่างหนึ่งเกินเกณฑ์)", "Wisconsin DOT Statewide Average Crash Rates and UCLs", "หน้า 9 Figure 3", "Segment Crash Rate > Crash Rate UCL OR KAB Crash Rate > KAB Crash Rate UCL", WI),
    ("ใช้ Severity ร่วมกับวิธีอื่น", NATT, "หน้า 55", "สามารถใช้วิธี Accident Severity ร่วมกับวิธีการอื่นได้ ในการศึกษาต่อๆไป", ""),
    ("L = 1 กม.", NATT, "หน้า 35–36 และหน้า 51", "สำหรับช่วงถนน 1 กิโลเมตร", ""),
]: ref.append(list(r))
for c in ref[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = blue
for col, w in zip("ABCDE", (34, 46, 30, 70, 40)): ref.column_dimensions[col].width = w
for row in ref.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    if row[4].value: row[4].hyperlink = row[4].value; row[4].font = Font(color="0563C1", underline="single")
wb.move_sheet("สรุปผล", offset=-wb.sheetnames.index("สรุปผล"))
wb.save(out); print("saved", out)
print(M.assign(df=M.R / M.Rc).nlargest(10, "df")[["id", "ถนน", "หน่วยงาน", "A", "V", "R", "Rc", "SI", "ระดับ_SI", "ระดับรวม (SI+Rate)"]].round(3).to_string())
