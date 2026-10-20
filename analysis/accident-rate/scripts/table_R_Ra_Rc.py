"""ตารางค่า R, Ra, Rc, DF ทั้ง 598 คลัสเตอร์ (ค่าตัวเลขจริง ไม่ใช่สูตร) + Ra รายกลุ่มตาม HSM สมการ 4-10"""
import os
import numpy as np, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "outputs", "ค่า-R-Ra-Rc-598คลัสเตอร์.xlsx")
M = pd.read_pickle(os.path.join(HERE, "data", "M598.pkl"))
T, L, K = 3, 1, 1.645

# ---------- Ra รายกลุ่ม: ตรวจว่า ΣA/Σm = Σ(V×R)/ΣV (HSM สมการ 4-10)
g = []
for ag, d in M.groupby("หน่วยงาน"):
    g.append({"หน่วยงาน (กลุ่มอ้างอิง)": ag, "จำนวนคลัสเตอร์": len(d), "ΣA (อุบัติเหตุ)": int(d.A.sum()),
              "ΣMEV (ล้านคัน)": d.m.sum(), "Ra = ΣA / ΣMEV": d.A.sum() / d.m.sum(),
              "ΣV (คัน/วัน)": d.V.sum(), "Σ(V × R)": (d.V * d.R).sum(), "Ra = Σ(V×R) / ΣV  [HSM 4-10]": (d.V * d.R).sum() / d.V.sum()})
G = pd.DataFrame(g)
assert np.allclose(G["Ra = ΣA / ΣMEV"], G["Ra = Σ(V×R) / ΣV  [HSM 4-10]"])

# ---------- ตารางหลัก
D = M.copy()
D["DF = R/Rc"] = D.R / D.Rc
D["ส่วน K×√(Ra/MEV)"] = K * np.sqrt(D.Ra / D.m)
D["ส่วน 1/(2MEV)"] = 1 / (2 * D.m)
order = {"high": 0, "medium": 1, "low": 2}
D = D.assign(_o=D["ระดับ_Rate"].map(order)).sort_values(["_o", "DF = R/Rc"], ascending=[True, False])
cols = {"id": "id", "ถนน": "ถนน", "จังหวัด": "จังหวัด", "หน่วยงาน": "หน่วยงาน", "A": "A (อุบัติเหตุ 3 ปี)",
        "AADT 2566": "AADT 2566", "AADT 2567": "AADT 2567", "AADT 2568": "AADT 2568", "V": "V = AADT เฉลี่ย 3 ปี",
        "m": "MEV = 365×T×V×L/10⁶", "R": "R (ครั้ง/ล้านคัน)", "Ra": "Ra ของกลุ่ม",
        "ส่วน K×√(Ra/MEV)": "K×√(Ra/MEV)", "ส่วน 1/(2MEV)": "1/(2MEV)", "Rc": "Rc", "DF = R/Rc": "DF = R/Rc",
        "ระดับ_Rate": "ระดับ Rate", "SI": "SI", "ระดับ_SI": "ระดับ SI", "ระดับรวม (SI+Rate)": "ระดับรวม"}
TBL = D[list(cols)].rename(columns=cols)
TBL.insert(0, "ลำดับ", range(1, len(TBL) + 1))
TBL["T"] = T; TBL["L"] = L

# ---------- สถิติรายกลุ่ม
S = (D.groupby("หน่วยงาน").agg(n=("id", "size"), R_ต่ำสุด=("R", "min"), R_มัธยฐาน=("R", "median"), R_สูงสุด=("R", "max"),
                                Rc_ต่ำสุด=("Rc", "min"), Rc_มัธยฐาน=("Rc", "median"), Rc_สูงสุด=("Rc", "max"),
                                สูง=("ระดับ_Rate", lambda s: (s == "high").sum()), ปานกลาง=("ระดับ_Rate", lambda s: (s == "medium").sum()),
                                ต่ำ=("ระดับ_Rate", lambda s: (s == "low").sum())).reset_index())

with pd.ExcelWriter(OUT, engine="openpyxl") as w:
    TBL.to_excel(w, sheet_name="R Ra Rc รายคลัสเตอร์", index=False)
    G.to_excel(w, sheet_name="Ra รายกลุ่ม (HSM 4-10)", index=False)
    S.to_excel(w, sheet_name="สถิติรายกลุ่ม", index=False)
    pd.DataFrame({"สูตร": [
        "R = A × 1,000,000 / (365 × T × V × L)  — ณัฐพงศ์ (2558) สมการ 3.2",
        "MEV = 365 × T × V × L / 1,000,000  (ล้านคันที่ผ่าน) — HSM สมการ 4-8 (L = 1 กม.)",
        "Ra = Σ(Vᵢ × Rᵢ) / ΣVᵢ = ΣA / ΣMEV  — AASHTO (2010) HSM สมการ 4-10 · คิดแยกตามหน่วยงาน",
        "Rc = Ra + K × [Ra / MEV]^0.5 + 1 / (2 × MEV)  — ณัฐพงศ์ (2558) สมการ 3.4 · HSM สมการ 4-11 · K = 1.645",
        "DF = R / Rc  — ณัฐพงศ์ (2558) สมการ 3.3",
        "ระดับ Rate: ต่ำ R ≤ Ra · ปานกลาง Ra < R ≤ Rc · สูง R > Rc (DF > 1)  — ลำดับขั้นตาม MnDOT (2010)",
        "T = 3 ปี (2566–2568) · L = 1 กม. · V = ค่าเฉลี่ย AADT 3 ปี",
        "ค่าในไฟล์นี้เป็นตัวเลขที่คำนวณแล้ว (ไม่ใช่สูตร) · ไฟล์ที่มีสูตร Excel ให้ตรวจ: ระดับจุดเสี่ยง-598คลัสเตอร์-SI+AccidentRate.xlsx"]}
    ).to_excel(w, sheet_name="สูตร", index=False)

wb = load_workbook(OUT); blue = PatternFill("solid", fgColor="1C5CAB"); orange = PatternFill("solid", fgColor="EB6834")
LV = {"high": "F8CBAD", "medium": "FFE699", "low": "C6EFCE"}
fmt = {"AADT": "#,##0", "V =": "#,##0", "MEV": "#,##0.00", "R (": "0.0000", "Ra": "0.0000", "K×": "0.0000", "1/(2MEV)": "0.00000",
       "Rc": "0.0000", "DF": "0.00", "SI": "0.00", "ΣMEV": "#,##0.00", "ΣV": "#,##0", "Σ(V": "#,##0.00", "R_": "0.0000"}
for ws in wb:
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = blue; c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 40; ws.freeze_panes = "B2"
    for j, c in enumerate(ws[1], 1):
        ws.column_dimensions[get_column_letter(j)].width = 14
        f = next((v for k, v in fmt.items() if str(c.value).startswith(k)), None)
        if f:
            for r in range(2, ws.max_row + 1): ws.cell(r, j).number_format = f
ws = wb["R Ra Rc รายคลัสเตอร์"]; h = {c.value: c.column for c in ws[1]}
for n in ("R (ครั้ง/ล้านคัน)", "Ra ของกลุ่ม", "Rc", "DF = R/Rc"): ws.cell(1, h[n]).fill = orange
ws.column_dimensions[get_column_letter(h["ถนน"])].width = 34; ws.auto_filter.ref = ws.dimensions
for r in range(2, ws.max_row + 1):
    for n in ("ระดับ Rate", "ระดับ SI", "ระดับรวม"):
        c = ws.cell(r, h[n]); c.fill = PatternFill("solid", fgColor=LV[c.value])
wb["Ra รายกลุ่ม (HSM 4-10)"].column_dimensions["A"].width = 22
wb["สูตร"].column_dimensions["A"].width = 120
wb.save(OUT); print("saved", os.path.abspath(OUT))
print(G.round(4).to_string(index=False)); print(S.round(4).to_string(index=False))
print(TBL.head(10)[["ลำดับ", "id", "ถนน", "หน่วยงาน", "A (อุบัติเหตุ 3 ปี)", "V = AADT เฉลี่ย 3 ปี", "MEV = 365×T×V×L/10⁶", "R (ครั้ง/ล้านคัน)", "Ra ของกลุ่ม", "Rc", "DF = R/Rc"]].round(4).to_string(index=False))
TBL.to_csv(os.path.join(HERE, "..", "outputs", "ค่า-R-Ra-Rc-598คลัสเตอร์.csv"), index=False, encoding="utf-8-sig")
