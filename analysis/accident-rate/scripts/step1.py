"""ขั้นที่ 1: ถนนกรมทางหลวง 14 ตอนควบคุม (69 คลัสเตอร์) ที่ไม่มีจุดสำรวจในตอนเอง
ใช้ 'สถานีตัวแทน' (station_id) ที่กรมทางหลวงกำหนดให้ตอนนั้นในชั้นข้อมูล hgds_aadt_line_2567 แล้วดึงค่าสถานีเดียวกันครบ 3 ปี"""
import os
import json, re, urllib.parse
import numpy as np, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
PROJ = os.environ.get("MTEC_OUT", os.path.join(HERE, "..", "outputs")) + "/"
DL = os.environ.get("MTEC_RAW", os.path.join(HERE, "..", "data-raw")) + "/"
SP = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data") + "/"
q14 = json.load(open(SP + "q14.json", encoding="utf-8"))
pnt = {x["station_id"]: x for x in json.load(open(SP + "q14b.json", encoding="utf-8"))["hgds_aadt_pnt_2567"]}
T = {y: pd.read_excel(DL + f"aadt-{y - 2500}.xlsx") for y in (2566, 2567, 2568)}
P68 = pd.read_csv(DL + "ข้อมูลปริมาณจราจรบนทางหลวงรายจุดสำรวจ (AADT_Pnt_2568).csv", encoding="utf-8-sig")

DOH_DS = "https://opendata.doh.go.th/dataset/ed101df4-f0a1-4d7f-b9d9-76859b2ca73e/resource/"
PAGE = {2566: DOH_DS + "6f55e818-7d88-4466-b345-89348882f6c9", 2567: DOH_DS + "f1a73a43-f496-4f35-a246-b0e5b373c3bd",
        2568: DOH_DS + "f88f9c4e-b32e-4cd1-bf51-4e1f84ba7b16"}
def wfs(layer, cql, props):
    return "https://onemap-layers.doh.go.th/geoserver/wfs?" + urllib.parse.urlencode(
        {"service": "WFS", "version": "2.0.0", "request": "GetFeature", "typeNames": "dohonemap:" + layer,
         "outputFormat": "application/json", "propertyName": props, "CQL_FILTER": cql})

def tab(y, h, km):
    d = T[y]; m = d[(pd.to_numeric(d["ทางหลวงสาย"], errors="coerce") == h) & (d["จุดสำรวจ"].astype(str).str.strip() == km)]
    v = sorted(set(m["รวม"])); assert len(v) == 1, (y, h, km, v); return int(v[0]), int(m["ตอนควบคุม"].iloc[0])

sec_rows = []
for key, r in q14.items():
    h, c = map(int, key.split("-"))
    L67 = r["hgds_aadt_line_2567"][0]; st = L67["station_id"]; s = pnt[st]
    sh, skm = int(s["highway_code"]), s["km"]
    v = {}; cs = {}
    for y in (2566, 2567, 2568): v[y], cs[y] = tab(y, sh, skm)
    assert v[2567] == L67["vehcat_aadt_all"] == s["vehcat_aadt_all"]                       # ตาราง opendata = GeoServer
    p68 = P68[(P68.highway_code.astype(int) == sh) & (P68.km.astype(str) == skm)].vehcat_aadt_all
    assert set(p68) == {v[2568]}                                                             # = ไฟล์ AADT_Pnt_2568 ที่ผู้ใช้โหลด
    same = (sh == h)
    sec_rows.append({
        "ทางหลวง": h, "ตอนควบคุม": f"{c:04d}", "ชื่อตอน": L67["road_name"], "กม. ตอน": f"{L67['km_start']} – {L67['km_end']}",
        "station_id ตัวแทน (ทล. กำหนด)": st, "สถานีตัวแทน: ทางหลวง": sh, "ตอน": s["road_code"], "ชื่อตอนของสถานี": s["road_name"], "กม. สถานี": skm,
        "lat สถานี": s["latitude"], "lon สถานี": s["longitude"], "ประเภทสถานี": s["station_type"],
        "AADT 2566": v[2566], "AADT 2567": v[2567], "AADT 2568": v[2568], "AADT เฉลี่ย 3 ปี": round((v[2566] + v[2567] + v[2568]) / 3, 1),
        "ลิงก์ ① สถานีตัวแทนของตอน (ปี 2567)": wfs("hgds_aadt_line_2567", f"highway_code='{h:04d}' AND road_code='{c:04d}'",
                                                  "highway_code,road_code,road_name,km_start,km_end,station_id,station_type,vehcat_aadt_all,traffic_year"),
        "ลิงก์ ② 2566 (opendata ทล.)": PAGE[2566], "วิธีค้น 2566": f"ค้นแถว ทางหลวงสาย {sh} · ตอนควบคุม {cs[2566]} · จุดสำรวจ {skm} → คอลัมน์ 'รวม'",
        "ลิงก์ ③ 2567 (opendata ทล.)": PAGE[2567], "วิธีค้น 2567": f"ค้นแถว ทางหลวงสาย {sh} · ตอนควบคุม {cs[2567]} · จุดสำรวจ {skm} → คอลัมน์ 'รวม'",
        "ลิงก์ ④ 2568 (opendata ทล.)": PAGE[2568], "วิธีค้น 2568": f"ค้นแถว ทางหลวงสาย {sh} · ตอนควบคุม {cs[2568]} · จุดสำรวจ {skm} → คอลัมน์ 'รวม'",
        "ลิงก์ ⑤ สถานีเดียวกัน 3 ปี (GeoServer)": wfs("hgds_aadt_pnt_2568", f"station_id={st}",
                                                   "highway_code,road_code,road_name,km,station_id,vehcat_aadt_all,traffic_year,latitude,longitude"),
        "หมายเหตุ": ("สถานีอยู่ทางหลวงสายเดียวกัน (ตอนถัดไป)" if same else
                    f"⚠ ทล. กำหนดสถานีของทางหลวง {sh} เป็นตัวแทนของตอนนี้ (ใช้ตามค่าอ้างอิงของ ทล.)") +
                   ("" if r["hgds_aadt_line_2565"][0]["station_id"] == st else f" · ชั้นข้อมูลปี 2565 เคยใช้สถานี {r['hgds_aadt_line_2565'][0]['station_id']}"),
    })
S = pd.DataFrame(sec_rows)

# ---------- คลัสเตอร์ใน 14 ตอน
k = pd.read_excel(PROJ + "AADT-ครบทุกคลัสเตอร์-630-3ปี.xlsx", sheet_name="คลัสเตอร์630")
def sec(s):
    m = re.search(r"ทางหลวง (\d+) ตอนควบคุม (\d+)", str(s)); return (int(m.group(1)), int(m.group(2))) if m else None
k["sec"] = k["ตอนที่จับคู่"].map(sec)
S["key"] = list(zip(S["ทางหลวง"], S["ตอนควบคุม"].astype(int)))
C = k[k.sec.isin(S.key)].merge(S[["key", "station_id ตัวแทน (ทล. กำหนด)", "AADT 2566", "AADT 2567", "AADT 2568"]], left_on="sec", right_on="key")
assert (C["AADT_2567"] == C["AADT 2567"]).all()
C = C[["id", "ถนน", "จังหวัด", "route_code", "กม. คลัสเตอร์", "ตอนที่จับคู่", "station_id ตัวแทน (ทล. กำหนด)", "A", "เสียชีวิต", "บาดเจ็บ", "SI", "ระดับ_SI",
       "AADT 2566", "AADT 2567", "AADT 2568"]].sort_values(["ตอนที่จับคู่", "id"]).reset_index(drop=True)
C.insert(15, "V = AADT เฉลี่ย 3 ปี", 0.0); C.insert(16, "T (ปี)", 3); C.insert(17, "L (กม.)", 1)
print("ตอน", len(S), "คลัสเตอร์", len(C), "อุบัติเหตุ", C.A.sum())

out = PROJ + "ขั้นที่1-AADT-กรมทางหลวง-14ตอน-3ปี.xlsx"
with pd.ExcelWriter(out, engine="openpyxl") as w:
    S.drop(columns="key").to_excel(w, sheet_name="14 ตอนควบคุม", index=False)
    C.to_excel(w, sheet_name="69 คลัสเตอร์ + Accident Rate", index=False)
    pd.DataFrame({"ขั้นตอนตรวจสอบ": [
        "ทำไมต้องใช้สถานีตัวแทน: 14 ตอนนี้ไม่มีจุดสำรวจในตอนเอง (ส่วนใหญ่ปี 2568 เป็นสถานีอัตโนมัติ AP ที่ไม่เผยแพร่ค่า) กรมทางหลวงจึงกำหนด station_id ของสถานีใกล้เคียงให้เป็นค่าของตอนในชั้นข้อมูล AADT รายตอน",
        "① เปิดลิงก์ ① ของตอนนั้น → ดู station_id และ vehcat_aadt_all (traffic_year 2024 = พ.ศ. 2567) เป็นหลักฐานว่า ทล. ใช้สถานีนี้แทนตอนนี้",
        "② ③ ④ เปิดหน้า opendata ปีนั้น → กดดาวน์โหลด/ดูตาราง → ค้นแถวตามคอลัมน์ 'วิธีค้น' → อ่านค่าในคอลัมน์ 'รวม' = AADT ของปีนั้น",
        "⑤ (ทางเลือก) ลิงก์ GeoServer แสดงค่าสถานีเดียวกันปี 2568 พร้อมพิกัด เปลี่ยนเลข 2568 ในลิงก์เป็น 2567 หรือ 2565 เพื่อดูปีอื่นได้ (ไม่มีชั้นข้อมูลปี 2566)",
        "ปี 2566 ใช้ตาราง opendata aadt-66 เท่านั้น เพราะ GeoServer/One Map ไม่มีชั้นปี 2566 (ไฟล์ที่ชื่อ AADT_Pnt_2566 บน One Map เป็นข้อมูลปี 2565 / traffic_year 2022)",
        "ตรวจแล้ว: ค่า 2567 ในตาราง opendata = ค่าในชั้นข้อมูล GeoServer ทุกสถานี และค่า 2568 = ไฟล์ AADT_Pnt_2568 ที่ดาวน์โหลดจาก One Map ทุกสถานี",
        "สูตร Accident Rate (สมการ 2.2): R = A × 1,000,000 / (365 × T × V × L) · V = AADT เฉลี่ย 2566–2568 · T = 3 ปี · L = 1 กม.",
    ]}).to_excel(w, sheet_name="วิธีตรวจสอบ", index=False)

wb = load_workbook(out); blue = PatternFill("solid", fgColor="1C5CAB"); orange = PatternFill("solid", fgColor="EB6834")
for ws in wb:
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = blue; c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 45; ws.freeze_panes = "B2"
ws = wb["14 ตอนควบคุม"]; h = {c.value: c.column for c in ws[1]}
for j in range(1, ws.max_column + 1): ws.column_dimensions[get_column_letter(j)].width = 14
for n in ("ชื่อตอน", "ชื่อตอนของสถานี", "หมายเหตุ"): ws.column_dimensions[get_column_letter(h[n])].width = 32
for n in [x for x in h if str(x).startswith("วิธีค้น")]: ws.column_dimensions[get_column_letter(h[n])].width = 44
for r in range(2, ws.max_row + 1):
    for n in [x for x in h if str(x).startswith("ลิงก์")]:
        c = ws.cell(r, h[n]); c.hyperlink = c.value; c.value = "เปิดลิงก์"; c.font = Font(color="0563C1", underline="single")
    for n in ("AADT 2566", "AADT 2567", "AADT 2568", "AADT เฉลี่ย 3 ปี"): ws.cell(r, h[n]).number_format = "#,##0"
    if "⚠" in str(ws.cell(r, h["หมายเหตุ"]).value): ws.cell(r, h["หมายเหตุ"]).fill = PatternFill("solid", fgColor="FFF2CC")

ws = wb["69 คลัสเตอร์ + Accident Rate"]; h = {c.value: c.column for c in ws[1]}; n = ws.max_row
L = lambda name: get_column_letter(h[name])
extra = ["ตัวตั้ง A×1,000,000", "ตัวหาร 365×T×V×L", "R (ครั้ง/ล้านคัน-กม.)", "อันดับ R ใน 69 วง"]
for j, name in enumerate(extra): ws.cell(1, ws.max_column + 1 if j else ws.max_column + 1, name)
h = {c.value: c.column for c in ws[1]}
for name in extra + ["V = AADT เฉลี่ย 3 ปี"]: ws.cell(1, h[name]).fill = orange
for r in range(2, n + 1):
    ws[f"{L('V = AADT เฉลี่ย 3 ปี')}{r}"] = f"=AVERAGE({L('AADT 2566')}{r}:{L('AADT 2568')}{r})"
    ws[f"{L(extra[0])}{r}"] = f"={L('A')}{r}*1000000"
    ws[f"{L(extra[1])}{r}"] = f"=365*{L('T (ปี)')}{r}*{L('V = AADT เฉลี่ย 3 ปี')}{r}*{L('L (กม.)')}{r}"
    ws[f"{L(extra[2])}{r}"] = f"={L(extra[0])}{r}/{L(extra[1])}{r}"
    ws[f"{L(extra[3])}{r}"] = f"=RANK({L(extra[2])}{r},${L(extra[2])}$2:${L(extra[2])}${n},0)"
    for name in ("AADT 2566", "AADT 2567", "AADT 2568", "V = AADT เฉลี่ย 3 ปี", extra[0], extra[1]): ws[f"{L(name)}{r}"].number_format = "#,##0"
    ws[f"{L(extra[2])}{r}"].number_format = "0.0000"
for j in range(1, ws.max_column + 1): ws.column_dimensions[get_column_letter(j)].width = 13
for name, wd in (("ถนน", 32), ("ตอนที่จับคู่", 26)): ws.column_dimensions[L(name)].width = wd
ws.auto_filter.ref = ws.dimensions
wb["วิธีตรวจสอบ"].column_dimensions["A"].width = 150
wb.save(out); print("saved", out)

C["V"] = C[["AADT 2566", "AADT 2567", "AADT 2568"]].mean(axis=1); C["R"] = C.A * 1e6 / (365 * 3 * C.V)
print(S[["ทางหลวง", "ตอนควบคุม", "ชื่อตอน", "station_id ตัวแทน (ทล. กำหนด)", "สถานีตัวแทน: ทางหลวง", "กม. สถานี", "AADT 2566", "AADT 2567", "AADT 2568", "หมายเหตุ"]].to_string())
print(C.groupby("ตอนที่จับคู่").agg(n=("id", "size"), A=("A", "sum")).to_string())
print(C.sort_values("R", ascending=False)[["id", "ถนน", "A", "V", "R"]].head(5).round(3).to_string())
